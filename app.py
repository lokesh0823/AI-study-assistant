import streamlit as st
import PyPDF2
import io
import docx
from groq import Groq
from fpdf import FPDF

client = Groq(api_key=st.secrets["GROQ_API_KEY"])

st.set_page_config(page_title="AI Study Assistant", page_icon="🧠")
st.title("🧠 AI Study Assistant")
st.write("Upload your notes and let AI help you study smarter.")
st.divider()

uploaded_file = st.file_uploader("Upload your notes (PDF, TXT or Word)", type=["pdf", "txt", "docx"])

def extract_text(file):
    if file.type == "application/pdf":
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file.read()))
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text
    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(io.BytesIO(file.read()))
        text = "\n".join([para.text for para in doc.paragraphs])
        return text
    else:
        return file.read().decode("utf-8")

def ask_groq(prompt):
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.choices[0].message.content

class PDF(FPDF):
    def clean_text(self, text):
        return ''.join(c if ord(c) < 128 else ' ' for c in text)

def generate_pdf(result, topic):
    pdf = PDF()
    pdf.add_page()
    pdf.set_margins(15, 15, 15)
    pdf.set_font("Helvetica", "B", 20)
    pdf.set_text_color(0, 100, 200)
    pdf.cell(0, 12, "Study Summary", ln=True, align="C")
    pdf.ln(4)
    pdf.set_font("Helvetica", "I", 11)
    pdf.set_text_color(100, 100, 100)
    pdf.multi_cell(0, 8, pdf.clean_text(f"Topic: {topic}"))
    pdf.ln(4)
    pdf.set_draw_color(0, 100, 200)
    pdf.line(15, pdf.get_y(), 195, pdf.get_y())
    pdf.ln(6)
    pdf.set_text_color(0, 0, 0)
    in_code_block = False
    for line in result.split("\n"):
        line = line.strip()
        if line.startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block:
            continue
        if not line:
            pdf.ln(2)
        elif line.startswith("#"):
            pdf.set_font("Helvetica", "B", 13)
            pdf.set_text_color(0, 100, 200)
            pdf.multi_cell(0, 8, pdf.clean_text(line.replace("#", "").strip()[:200]))
            pdf.set_text_color(0, 0, 0)
        elif line.startswith("*") or line.startswith("-"):
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 7, pdf.clean_text("  - " + line[1:].strip()[:200]))
        elif line.startswith("**") and line.endswith("**"):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 7, pdf.clean_text(line.replace("**", "").strip()[:200]))
        else:
            pdf.set_font("Helvetica", "", 11)
            pdf.multi_cell(0, 7, pdf.clean_text(line[:200]))
    return bytes(pdf.output())

if uploaded_file:
    text = extract_text(uploaded_file)
    st.success(f"✅ File uploaded! ({len(text)} characters extracted)")

    with st.spinner("Detecting topic..."):
        try:
            topic = ask_groq(f"In one short sentence, what subject or topic are these notes about?\n\n{text[:2000]}")
            st.info("📚 Topic detected: " + topic.replace("*", "").replace("_", ""))
        except Exception as e:
            topic = "Unknown"

    custom_instruction = st.text_input(
        "Optional: Add custom instructions (e.g. 'summarize like I'm a beginner' or 'focus only on key concepts')",
        placeholder="Leave empty for default behavior"
    )

    st.divider()

    tab1, tab2, tab3 = st.tabs(["📝 Summarize", "❓ Quiz Me", "💬 Ask a Question"])

    with tab1:
        if st.button("Generate Summary", use_container_width=True):
            with st.spinner("Summarizing your notes..."):
                try:
                    result = ask_groq(
                        f"Summarize these notes in detail. For each concept: "
                        f"1) Explain the logic clearly in plain English "
                        f"2) If there is code, show it and explain line by line what it does "
                        f"3) Give a real world analogy where possible. "
                        f"Be thorough, not basic. "
                        f"{'Additional instructions: ' + custom_instruction if custom_instruction else ''}\n\n{text[:10000]}"
                    )
                    st.markdown(result)
                    st.session_state['summary_result'] = result
                    st.session_state['summary_topic'] = topic
                except Exception as e:
                    st.error(f"Error: {str(e)}")


    with tab2:
        num_questions = st.slider("Number of questions", 3, 10, 5)
        difficulty = st.selectbox("Difficulty", ["Easy", "Medium", "Hard"])
        if st.button("Generate Quiz", use_container_width=True):
            with st.spinner("Generating quiz..."):
                try:
                    result = ask_groq(
                        f"Generate {num_questions} {difficulty} difficulty quiz questions with answers. "
                        f"{'Additional instructions: ' + custom_instruction if custom_instruction else ''}"
                        f"Format each as Q: question and A: answer:\n\n{text[:10000]}"
                    )
                    st.markdown(result)
                except Exception as e:
                    st.error(f"Error: {str(e)}")

    with tab3:
        if "messages" not in st.session_state:
            st.session_state.messages = []

        for msg in st.session_state.messages:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])

        question = st.chat_input("Ask anything about your notes...")
        if question:
            st.session_state.messages.append({"role": "user", "content": question})
            with st.chat_message("user"):
                st.markdown(question)

            with st.chat_message("assistant"):
                with st.spinner("Thinking..."):
                    try:
                        history = "\n".join([f"{m['role']}: {m['content']}" for m in st.session_state.messages[-6:]])
                        answer = ask_groq(
                            f"Based on these notes:\n\n{text[:8000]}\n\n"
                            f"Conversation so far:\n{history}\n\n"
                            f"Answer the latest question thoroughly."
                        )
                        st.markdown(answer)
                        st.session_state.messages.append({"role": "assistant", "content": answer})
                    except Exception as e:
                        st.error(f"Error: {str(e)}")